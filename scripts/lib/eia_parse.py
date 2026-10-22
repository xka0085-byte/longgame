#!/usr/bin/env python3
"""Parse an EIA-860M generator xlsx and print a compact JSON summary.
Usage: python eia_parse.py <file.xlsx>
Only reads the workbook, no network, no writes.

硬约束：必需的工作表或列缺失、MW 单元格无法解析时，不得回退成 0.0，
必须非零退出并输出 errors，让采集与发布链路中断。
"""
import sys, json
import openpyxl

path = sys.argv[1]
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
out = {"sheets": {}, "plannedMW": 0.0, "canceledMW": 0.0, "plannedByYear": {}, "retiredMW": 0.0}
errors = []

REQUIRED_SHEETS = ["Operating", "Planned", "Retired", "Canceled or Postponed"]
for name in REQUIRED_SHEETS:
    if name not in wb.sheetnames:
        errors.append(f"missing sheet: {name}")

for name in wb.sheetnames:
    ws = wb[name]
    out["sheets"][name] = max(ws.max_row - 3, 0)

def col_index(ws, sheet):
    hdr = [c.value for c in next(ws.iter_rows(min_row=3, max_row=3))]
    idx = {h: i for i, h in enumerate(hdr) if h}
    return idx

def require(idx, col, sheet):
    if col not in idx:
        errors.append(f"missing column in {sheet}: {col}")
        return None
    return idx[col]

def num(v):
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        errors.append(f"non-numeric MW value: {v!r}")
        return None

if "Planned" in wb.sheetnames:
    ws = wb["Planned"]
    idx = col_index(ws, "Planned")
    ci = require(idx, "Nameplate Capacity (MW)", "Planned")
    yi = require(idx, "Planned Operation Year", "Planned")
    for row in ws.iter_rows(min_row=4, values_only=True):
        mw = num(row[ci]) if ci is not None else None
        if mw is not None:
            out["plannedMW"] += mw
        y = row[yi] if yi is not None else None
        if y not in (None, ""):
            key = str(int(y)) if isinstance(y, (int, float)) else str(y)
            out["plannedByYear"][key] = out["plannedByYear"].get(key, 0) + 1

if "Canceled or Postponed" in wb.sheetnames:
    ws = wb["Canceled or Postponed"]
    idx = col_index(ws, "Canceled or Postponed")
    ci = require(idx, "Nameplate Capacity (MW)", "Canceled or Postponed")
    for row in ws.iter_rows(min_row=4, values_only=True):
        mw = num(row[ci]) if ci is not None else None
        if mw is not None:
            out["canceledMW"] += mw

if "Retired" in wb.sheetnames:
    ws = wb["Retired"]
    idx = col_index(ws, "Retired")
    ci = require(idx, "Nameplate Capacity (MW)", "Retired")
    for row in ws.iter_rows(min_row=4, values_only=True):
        mw = num(row[ci]) if ci is not None else None
        if mw is not None:
            out["retiredMW"] += mw

out["plannedMW"] = round(out["plannedMW"], 1)
out["canceledMW"] = round(out["canceledMW"], 1)
out["retiredMW"] = round(out["retiredMW"], 1)

if errors:
    out["errors"] = errors
    print(json.dumps(out, ensure_ascii=False))
    print("EIA parse failed: " + "; ".join(errors[:6]), file=sys.stderr)
    sys.exit(1)

print(json.dumps(out))
